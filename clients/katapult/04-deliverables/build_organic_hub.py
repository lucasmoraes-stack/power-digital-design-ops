"""Rebuild the Organic Social part of the Katapult Deliverables Hub from finished PNGs (and reel MP4s).

usage: python build_organic_hub.py <hub_in.html> <hub_out.html> <files_dir>

- Replaces the old organic markup (`.q4-bar` + `<main data-ch="organic">`) with `<!-- ORG:START -->` image-based posts,
  one section per post per month, and rewrites the organic index (month switch, November calendar, post cards).
- Leaves every banner block (K19 to K22, KB, KN) untouched; patches the NAV router only where it reads organic posts.
- Writes the published images to <files_dir>/organic/<month>/<post>/<name>.jpg (+ t/ thumbnails) and the reels.
  Publish them with the hub page as `files` (paths relative to the page).
"""
import html, json, pathlib, re, shutil, sys
from datetime import date
from PIL import Image

ROOT = pathlib.Path(__file__).resolve().parent            # clients/katapult/04-deliverables
WORK = ROOT.parent / "03-work" / "social"
FIGMA = "https://www.figma.com/design/D8FpNuKI3uPGjCPqonIn3C/Progressive-Global---LAB?node-id="
E = lambda s: html.escape(str(s), quote=True)

# ---------------------------------------------------------------- months
OCT_THEME = {**{n: "Educational" for n in ("01", "02", "03")}, **{n: "Inspirational" for n in ("04", "05", "06")},
             **{n: "Engaging" for n in ("07", "08", "09")}, **{n: "Promotional" for n in ("10", "11", "12")},
             **{n: "Brand storytelling" for n in ("13", "14", "15")}}

def month_oct():
    d = ROOT / "social" / "2026-10-organic"
    posts = []
    for p in json.loads((d / "frames.json").read_text(encoding="utf-8")):
        fmt = p["format"]
        posts.append({
            "num": p["post"], "title": p["title"], "theme": OCT_THEME[p["post"]], "format": fmt, "date": None,
            "slides": [{"png": d / f["file"], "w": f["w"], "h": f["h"]} for f in p["frames"]],
            "reel": None, "caption": None,
            "links": [("Open in Figma", FIGMA + p["section_id"].replace(":", "-"))],
            "notes": [],
        })
    return {"id": "2026-10", "label": "October 2026", "short": "October", "status": "Final",
            "status_note": "Exported from Figma (page kat os) 2026-10-16",
            "lede": "The final set, finished in Figma. These are the files that go to the client.",
            "posts": posts}

NOV_NOTES = {
    "01": ["The sheet's “1 —” separators became numerals in discs (no dashes on the art)."],
    "02": ["“-->” is drawn as an arrow icon.", "“FInd” fixed to “Find”; the sheet labels two slides “Slide 4”, built as five slides."],
    "03": ["The art only reserves an empty zone (about 880 x 760, centred) for Instagram's native poll/quiz sticker; the options are not drawn.",
           "Sticker options from the sheet: New appliances 🧺 · Furniture 🛋️ · Mattress 🛏️ · Electronics 📷 · Tires 🛞.",
           "Instagram's poll and quiz stickers take up to 4 options; the sheet has 5. Client to pick which one to drop (or use a question sticker)."],
    "05": ["The sheet links a Shutterstock photo (2670393107) we don't hold. Built with a stand-in; swap if the client licenses it."],
    "06": ["Legal: “with a little interest” describes a lease as interest-bearing. Needs client legal and the reviewer's permission to use the name."],
    "07": ["Sheet typo “How's to Use” fixed to “How to Use”.", "Checkout, phone-number and order-review screens are drawn generic UI (no retailer, no real numbers). Swap for real recordings if the client has them."],
    "10": ["“feature product 1 to 5” had no files: built with the Studio product cutouts."],
    "13": ["“moving night” set as “movie night” (the footnote is couches, TVs and a night in). Confirm."],
    "15": ["“product 1 to 5” had no files: built with the Studio product cutouts."],
}

def month_nov():
    import openpyxl
    xl = ROOT.parent / "03-work" / "[KATAPULT] OS Content Calendar Outline (2026-11).xlsx"
    ws = openpyxl.load_workbook(xl)["November 2026"]
    d = ROOT / "social" / "2026-11-organic"
    posts = []
    for r in range(5, 20):
        g = lambda c: ws[f"{c}{r}"].value
        num = f"{int(g('A')):02d}"
        title = re.match(r"\[(.+?)\]", str(g("H"))).group(1).strip()
        fmt = str(g("D")).replace(" (1080x1350)", " 4:5").replace(" (1080x1920)", " 9:16")
        folder = next((x for x in (d / "png").glob(f"{num}-*") if x.is_dir()), None) if (d / "png").exists() else None
        slides = []
        if folder:
            for f in sorted(folder.glob("*.png")):
                w, h = Image.open(f).size
                slides.append({"png": f, "w": w, "h": h})
        reel = next(iter(sorted((d / "reels").glob(f"{num}-*.mp4"))), None) if (d / "reels").exists() else None
        poster = reel.with_suffix(".png") if reel and reel.with_suffix(".png").exists() else None
        notes = list(NOV_NOTES.get(num, []))
        if "Reel" in fmt and slides:
            # Reels ship as MP4 plus the storyboard (Lucas, 2026-10-16): video first, then one frame per beat
            fmt = fmt.replace("Reel 9:16", "Reel 9:16 storyboard")
            sb = folder / "storyboard.md"
            if sb.exists():
                rows = []
                for l in sb.read_text(encoding="utf-8").splitlines():
                    cells = [c.strip().replace("**", "") for c in l.strip().strip("|").split("|")]
                    if l.startswith("|") and len(cells) >= 4 and cells[0][:1].isdigit():
                        rows.append(f"Frame {cells[0]} · {cells[1]} s · {cells[2]} On screen: {cells[3]}")
                notes = ["Delivered as the edited MP4 (first) plus the storyboard: one frame per beat."] + rows + notes
        cap = g("K")
        posts.append({
            "num": num, "title": title, "theme": str(g("C")).split("/")[0].replace("Engaging", "Engaging").strip(),
            "format": fmt, "date": g("B").date(), "slides": slides,
            "reel": {"mp4": reel, "poster": poster} if reel else None,
            "caption": None if not cap or str(cap).strip() == "n/a" else str(cap).strip(),
            "links": [("Inspo", g("L").split()[0])] if g("L") else [],
            "notes": notes,
        })
    return {"id": "2026-11", "label": "November 2026", "short": "November", "status": "For review",
            "status_note": "Built 2026-10-16 from the client sheet",
            "lede": "Posting Nov 3 to Nov 30, built in the look of the October finals. For review before they go to Figma.",
            "posts": posts}

# ---------------------------------------------------------------- files
def stage(months, files_dir):
    out = pathlib.Path(files_dir)
    if (out / "organic").exists():
        shutil.rmtree(out / "organic")
    for m in months:
        for p in m["posts"]:
            for i, s in enumerate(p["slides"]):
                rel = f"organic/{m['id']}/{s['png'].parent.name}/{s['png'].stem}.jpg"
                trel = f"organic/{m['id']}/{s['png'].parent.name}/t/{s['png'].stem}.jpg"
                im = Image.open(s["png"]).convert("RGB")
                (out / rel).parent.mkdir(parents=True, exist_ok=True)
                (out / trel).parent.mkdir(parents=True, exist_ok=True)
                im.save(out / rel, "JPEG", quality=90, optimize=True, progressive=True)
                im.resize((432, round(432 * s["h"] / s["w"])), Image.LANCZOS).save(out / trel, "JPEG", quality=84, optimize=True)
                s["src"], s["thumb"] = rel, trel
            if p["reel"]:
                rel = f"organic/{m['id']}/reels/{p['reel']['mp4'].name}"
                (out / rel).parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(p["reel"]["mp4"], out / rel)
                p["reel"]["src"] = rel
                if p["reel"]["poster"]:
                    im = Image.open(p["reel"]["poster"]).convert("RGB")
                    prel = rel[:-4] + ".jpg"; trel = rel[:-4] + "-t.jpg"
                    im.save(out / prel, "JPEG", quality=88, optimize=True)
                    im.resize((432, round(432 * im.height / im.width)), Image.LANCZOS).save(out / trel, "JPEG", quality=84)
                    p["reel"]["poster_src"], p["reel"]["thumb"] = prel, trel

def card_thumb(p):
    if p["slides"]:
        return p["slides"][0]["thumb"], p["slides"][0]["h"] > p["slides"][0]["w"] * 1.5
    if p["reel"] and p["reel"].get("thumb"):
        return p["reel"]["thumb"], True
    return None, False

# ---------------------------------------------------------------- markup
def fmt_date(d):
    return d.strftime("%a %b ") + str(d.day) if d else ""

def post_section(m, p):
    pid = f"m{m['id'].replace('-', '')}-{p['num']}"
    n = len(p["slides"])
    meta = [("Month", m["label"])]
    if p["date"]: meta.append(("Date", fmt_date(p["date"])))
    meta += [("Theme", p["theme"]), ("Format", p["format"])]
    if n > 1: meta.append(("Slides", str(n)))
    dl = "".join(f"<div><dt>{E(k)}</dt><dd>{E(v)}</dd></div>" for k, v in meta)
    links = "".join(f'<a class="og-link" href="{E(u)}" target="_blank" rel="noopener">{E(t)} ↗</a>' for t, u in p["links"])
    figs = []
    if p["reel"]:
        r = p["reel"]
        poster = f' poster="{E(r["poster_src"])}"' if r.get("poster_src") else ""
        figs.append(f'<figure class="og-slide og-reel"><video controls playsinline preload="metadata"{poster} width="1080" height="1920">'
                    f'<source src="{E(r["src"])}" type="video/mp4"></video>'
                    f'<figcaption><span>Reel · MP4</span><a href="{E(r["src"])}" download>Download</a></figcaption></figure>')
    for i, s in enumerate(p["slides"]):
        tall = " og-tall" if s["h"] > s["w"] * 1.5 else ""
        word = "Frame" if "storyboard" in p["format"] else "Slide"
        lab = f"{word} {i + 1}" if n > 1 else ("Story" if tall else "Static")
        figs.append(f'<figure class="og-slide{tall}"><a href="{E(s["src"])}" target="_blank" rel="noopener">'
                    f'<img src="{E(s["src"])}" alt="{E(p["title"])}, {E(lab.lower())}" width="{s["w"]}" height="{s["h"]}" loading="lazy"></a>'
                    f'<figcaption><span>{E(lab)}</span><a href="{E(s["src"])}" download>Download</a></figcaption></figure>')
    if not figs:
        figs.append('<p class="og-empty">In production.</p>')
    side = ""
    if p["caption"]:
        side += f'<div class="og-box"><h3>Caption (client)</h3><p class="og-cap">{E(p["caption"])}</p></div>'
    if p["notes"]:
        side += '<div class="og-box og-flag"><h3>Notes</h3><ul>' + "".join(f"<li>{E(x)}</li>" for x in p["notes"]) + "</ul></div>"
    attrs = (f'data-month="{m["id"]}" data-num="{p["num"]}" data-title="{E(p["title"])}" data-theme="{E(p["theme"])}" '
             f'data-format="{E(p["format"])}" data-slides="{n}" data-date="{p["date"].isoformat() if p["date"] else ""}"')
    th, tall = card_thumb(p)
    attrs += f' data-thumb="{E(th or "")}" data-tall="{1 if tall else 0}" data-reel="{1 if p["reel"] else 0}"'
    return (f'<section class="og-post" id="{pid}" {attrs} aria-labelledby="{pid}-h">\n'
            f'  <div class="og-head"><span class="og-num">{p["num"]}</span><div class="og-titles"><span class="og-kicker">{E(m["label"])} · {E(m["status"])}</span>'
            f'<h2 id="{pid}-h">{E(p["title"])}</h2></div><div class="og-links">{links}</div></div>\n'
            f'  <dl class="og-meta">{dl}</dl>\n'
            f'  <div class="og-strip">{"".join(figs)}</div>\n'
            + (f'  <div class="og-side">{side}</div>\n' if side else "")
            + "</section>")

ORG_CSS = r"""
/* ORG:START */
/* ============ Organic social: months, calendar, image posts ============ */
.og-months{display:flex;flex-wrap:wrap;gap:4px;padding:4px;margin:0 0 28px;border:1px solid var(--rule-strong);border-radius:12px;width:fit-content;max-width:100%;background:var(--panel)}
.og-mbtn{font:inherit;cursor:pointer;display:flex;align-items:center;gap:10px;padding:8px 14px;border-radius:9px;border:0;background:transparent;color:var(--ink-2);min-width:0}
.og-mbtn b{font-size:15px;font-weight:500;color:var(--ink)}
.og-mbtn small{font-size:12.5px}
.og-mbtn[aria-pressed="true"]{background:var(--k-pink);color:var(--k-dark-blue)}
.og-mbtn[aria-pressed="true"] b{color:var(--k-dark-blue)}
.og-mbtn[aria-pressed="true"] .og-dot{background:var(--k-dark-blue)}
.og-mbtn:not([aria-pressed="true"]):hover b{color:var(--accent)}
.og-dot{width:8px;height:8px;border-radius:50%;background:var(--k-pink);flex:none}
.og-dot.is-final{background:var(--k-light-blue)}
.og-mpanel[hidden]{display:none}
.og-mhead{display:flex;flex-direction:column;gap:6px;margin-bottom:26px;max-width:80ch}
.og-mhead > div{display:flex;flex-wrap:wrap;align-items:center;gap:8px 12px}
.og-mhead h3{font-size:24px;font-weight:500;letter-spacing:-.01em;margin:0}
.og-mhead p{color:var(--ink-2);font-size:15px;margin:0}
.og-mhead .og-facts{font-size:13.5px;color:var(--ink-2)}
.og-mhead .og-facts b{color:var(--ink);font-weight:600}
.og-status{font-size:11px;font-weight:700;letter-spacing:.07em;text-transform:uppercase;padding:3px 9px;border-radius:999px;background:var(--k-pink);color:var(--k-dark-blue)}
.og-status.is-final{background:var(--k-light-blue)}
.og-week{margin-bottom:30px}
.og-week > h4{display:flex;align-items:baseline;gap:10px;margin:0 0 12px;font-size:12px;font-weight:700;letter-spacing:.09em;text-transform:uppercase;color:var(--ink-2)}
.og-week > h4 span{font-weight:500;letter-spacing:0;text-transform:none;font-size:13px}
.og-week > h4::after{content:"";flex:1;height:1px;background:var(--rule)}
.og-grid{display:grid;gap:16px;grid-template-columns:repeat(auto-fill,minmax(min(100%,200px),1fr))}
.og-card{display:flex;flex-direction:column;text-decoration:none;color:inherit;border-radius:12px;overflow:hidden;background:var(--panel);border:1px solid var(--rule);transition:border-color .15s,transform .15s}
.og-card:hover{border-color:var(--accent);transform:translateY(-2px)}
.og-card .og-th{position:relative;aspect-ratio:4/5;background:var(--code-bg);display:flex;justify-content:center;overflow:hidden}
.og-card .og-th img{height:100%;width:auto;max-width:100%;object-fit:cover;display:block}
.og-card .og-th .og-badge{position:absolute;right:8px;top:8px;font-size:11px;font-weight:700;padding:3px 8px;border-radius:999px;background:rgba(19,21,64,.85);color:#fff}
.og-card .og-cb{padding:10px 12px 12px;display:flex;flex-direction:column;gap:3px}
.og-card .og-when{font-size:12px;color:var(--ink-2);font-family:var(--font-mono)}
.og-card .og-t{font-size:15px;font-weight:500;line-height:1.25}
.og-card:hover .og-t{color:var(--accent)}
.og-card .og-f{font-size:12.5px;color:var(--ink-2)}
.og-lay{display:grid;grid-template-columns:minmax(0,360px) minmax(0,1fr);gap:32px;align-items:start}
.og-feed{position:sticky;top:16px;border:1px solid var(--rule-strong);border-radius:18px;padding:16px;background:var(--panel)}
.og-feed h4{margin:0 0 4px;font-size:12px;font-weight:700;letter-spacing:.09em;text-transform:uppercase;color:var(--ink-2)}
.og-feed p{margin:0 0 12px;font-size:12.5px;color:var(--ink-2);line-height:1.4}
.og-fgrid{display:grid;grid-template-columns:repeat(3,1fr);gap:2px;max-height:calc(100vh - 160px);overflow-y:auto;border-radius:6px}
.og-ft{position:relative;display:block;aspect-ratio:3/4;overflow:hidden;background:var(--code-bg)}
.og-ft img{width:100%;height:100%;object-fit:cover;display:block}
.og-ft.is-old{opacity:.4}
.og-ft:hover{opacity:1;outline:2px solid var(--accent);outline-offset:-2px}
.og-fi{position:absolute;top:5px;right:5px;line-height:0;filter:drop-shadow(0 1px 2px rgba(0,0,0,.5))}
@media (max-width:900px){.og-lay{grid-template-columns:1fr}.og-feed{position:static}.og-fgrid{max-height:none}}
@media (max-width:600px){
  .og-grid{grid-template-columns:repeat(2,minmax(0,1fr));gap:10px}
  .og-card .og-t{font-size:14px}
  .q4-mast .q4-wrap{padding-top:96px}
}

main[data-ch="organic"]{padding-bottom:64px}
.og-post{max-width:1400px;width:100%;min-width:0;box-sizing:border-box;margin-inline:auto;padding:18px max(16px,var(--gutter)) 0}
.og-post .hub-crumb{padding:0 0 18px;max-width:none}
.og-head{display:flex;flex-wrap:wrap;align-items:center;gap:12px 16px}
.og-num{font-family:var(--font-mono);font-size:14px;border:1px solid var(--rule-strong);border-radius:8px;padding:4px 9px;color:var(--ink-2)}
.og-titles{display:flex;flex-direction:column;gap:2px}
.og-kicker{font-size:12px;font-weight:700;letter-spacing:.08em;text-transform:uppercase;color:var(--ink-2)}
.og-head h2{font-size:clamp(24px,3vw,32px);font-weight:500;letter-spacing:-.01em;margin:0}
.og-links{margin-left:auto;display:flex;gap:8px;flex-wrap:wrap}
.og-link{font-size:13px;font-weight:700;text-decoration:none;border:1px solid var(--rule-strong);border-radius:999px;padding:5px 12px}
.og-link:hover{border-color:var(--accent)}
.og-meta{display:flex;flex-wrap:wrap;gap:6px 22px;margin:14px 0 18px}
.og-meta dt{font-size:11px;font-weight:700;letter-spacing:.08em;text-transform:uppercase;color:var(--ink-2)}
.og-meta dd{margin:0;font-size:15px}
.og-strip{max-width:100%;min-width:0;display:flex;gap:16px;overflow-x:auto;padding-bottom:14px;scroll-snap-type:x proximity}
.og-slide{flex:none;width:min(432px,82vw);margin:0;scroll-snap-align:start}
.og-slide.og-tall,.og-slide.og-reel{width:min(360px,74vw)}
.og-slide img,.og-slide video{display:block;width:100%;height:auto;border-radius:6px;box-shadow:0 0 0 1px var(--rule);background:var(--code-bg)}
.og-slide figcaption{display:flex;justify-content:space-between;font-size:13px;color:var(--ink-2);padding-top:6px}
.og-slide figcaption a{font-weight:700;color:var(--ink);text-decoration:none}
.og-slide figcaption a:hover{color:var(--accent)}
.og-empty{color:var(--ink-2);padding:40px 0}
.og-side{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,320px),1fr));gap:14px;margin-top:10px}
.og-box{border:1px solid var(--rule);border-radius:12px;padding:14px 16px;background:var(--panel)}
.og-box h3{font-size:12px;font-weight:700;letter-spacing:.08em;text-transform:uppercase;color:var(--ink-2);margin:0 0 8px}
.og-cap{white-space:pre-line;font-size:14.5px;line-height:1.5;margin:0}
.og-box ul{margin:0;padding-left:18px;font-size:14.5px;line-height:1.5}
.og-flag{border-color:var(--k-pink)}
/* ORG:END */
"""

ORG_JS = r"""
  function buildPosts(){
    var wrap=document.getElementById('hub-posts'); if(!wrap) return;
    var months={}, order=[];
    POSTS.forEach(function(s){ var m=s.dataset.month; if(!months[m]){ months[m]=[]; order.push(m); } months[m].push(s); });
    var META=ORG_MONTHS;
    order.sort().reverse();
    wrap.replaceChildren();
    var bar=el('div','og-months'); bar.setAttribute('role','group'); bar.setAttribute('aria-label','Month');
    wrap.appendChild(bar);
    function fmtDay(iso,opt){ return new Date(iso+'T12:00:00').toLocaleDateString('en-US',opt); }
    function card(s){
      var a=el('a','og-card'); a.href='#'+s.id;
      var th=el('div','og-th');
      if(s.dataset.thumb){ var im=el('img'); im.src=s.dataset.thumb; im.alt=''; im.decoding='async'; th.appendChild(im); }
      var n=+s.dataset.slides, f=s.dataset.format;
      if(/storyboard/.test(f)) th.appendChild(el('span','og-badge','Reel MP4 + '+n+' frames'));
      else if(s.dataset.reel==='1') th.appendChild(el('span','og-badge','Reel'));
      else if(n>1) th.appendChild(el('span','og-badge',n+' slides'));
      a.appendChild(th);
      var b=el('div','og-cb');
      b.appendChild(el('span','og-when')).textContent=s.dataset.num+(s.dataset.date?' · '+fmtDay(s.dataset.date,{weekday:'short',month:'short',day:'numeric'}):'');
      b.appendChild(el('span','og-t')).textContent=s.dataset.title;
      b.appendChild(el('span','og-f')).textContent=f.replace(' Post','').replace(/ (4:5|9:16)$/,'')+' · '+s.dataset.theme;
      a.appendChild(b); return a;
    }
    order.forEach(function(mid){
      var mm=META[mid]||{label:mid,status:''}, list=months[mid], fin=mm.status==='Final';
      var b=el('button','og-mbtn'); b.type='button'; b.dataset.month=mid;
      b.innerHTML='<i class="og-dot'+(fin?' is-final':'')+'"></i><b>'+mm.label+'</b><small>'+mm.status+'</small>';
      b.addEventListener('click',function(){ location.hash='#organic-'+mid; });
      bar.appendChild(b);
      var pn=el('div','og-mpanel'); pn.dataset.month=mid;
      var cnt={slides:0,car:0,stat:0,story:0,reel:0};
      list.forEach(function(s){ var f=s.dataset.format; cnt.slides+=+s.dataset.slides;
        if(/Carousel/.test(f)) cnt.car++; else if(/Story/.test(f)) cnt.story++; else if(/Reel/.test(f)) cnt.reel++; else cnt.stat++; });
      var facts=[[list.length,'posts'],[cnt.car,'carousels'],[cnt.stat,'statics'],[cnt.story,'Stories'],[cnt.reel,'Reels'],[cnt.slides,'frames']]
        .filter(function(c){return c[0];}).map(function(c){return '<b>'+c[0]+'</b> '+c[1];}).join(' · ');
      var hd=el('div','og-mhead'), row=el('div');
      row.appendChild(el('h3',null,mm.label));
      row.appendChild(el('span','og-status'+(fin?' is-final':''),mm.status));
      hd.appendChild(row);
      hd.appendChild(el('p',null,mm.lede));
      hd.appendChild(el('span','og-facts',facts+' · '+mm.note));
      pn.appendChild(hd);
      /* feed preview: the profile grid once the month is live, newest first, earlier months dimmed below */
      var lay=el('div','og-lay'), side=el('aside','og-feed'), mainc=el('div','og-main');
      side.appendChild(el('h4',null,'Feed preview'));
      side.appendChild(el('p',null,'The profile grid once every '+mm.label.split(' ')[0]+' post is live, newest first. Stories stay out of the grid; tiles crop to 3:4 like Instagram.'));
      var grid=el('div','og-fgrid');
      function feedOrder(arr){ return arr.filter(function(s){ return !/Story/.test(s.dataset.format); })
        .sort(function(x,y){ var a=x.dataset.date||x.dataset.num, c=y.dataset.date||y.dataset.num; return a<c?1:-1; }); }
      function tile(s,old){
        var a=el('a','og-ft'+(old?' is-old':'')); a.href='#'+s.id; a.title=s.dataset.num+' · '+s.dataset.title;
        if(s.dataset.thumb){ var im=el('img'); im.src=s.dataset.thumb; im.alt=s.dataset.title; im.decoding='async'; a.appendChild(im); }
        var f=s.dataset.format;
        if(/Reel/.test(f)) a.appendChild(el('i','og-fi','<svg viewBox="0 0 16 16" width="14" height="14"><path d="M4 2.5v11l9-5.5z" fill="#fff"/></svg>'));
        else if(+s.dataset.slides>1) a.appendChild(el('i','og-fi','<svg viewBox="0 0 16 16" width="14" height="14"><rect x="4.5" y="1.5" width="10" height="10" rx="1.5" fill="none" stroke="#fff" stroke-width="1.6"/><rect x="1.5" y="4.5" width="10" height="10" rx="1.5" fill="#fff"/></svg>'));
        return a;
      }
      feedOrder(list).forEach(function(s){ grid.appendChild(tile(s,false)); });
      order.filter(function(m){return m<mid;}).forEach(function(m){ feedOrder(months[m]).forEach(function(s){ grid.appendChild(tile(s,true)); }); });
      side.appendChild(grid);
      lay.appendChild(side); lay.appendChild(mainc); pn.appendChild(lay);
      var pnOuter=pn; pn=mainc;
      var dated=list.filter(function(s){return s.dataset.date;});
      if(dated.length===list.length){
        /* group by posting week, Monday first */
        var weeks={}, wk=[];
        list.slice().sort(function(x,y){return x.dataset.date<y.dataset.date?-1:1;}).forEach(function(s){
          var d=new Date(s.dataset.date+'T12:00:00'), mon=new Date(d); mon.setDate(d.getDate()-((d.getDay()+6)%7));
          var k=mon.getFullYear()+'-'+String(mon.getMonth()+1).padStart(2,'0')+'-'+String(mon.getDate()).padStart(2,'0');
          if(!weeks[k]){weeks[k]=[];wk.push(k);} weeks[k].push(s); });
        wk.forEach(function(k,i){
          var sec=el('section','og-week'), sun=new Date(k+'T12:00:00'); sun.setDate(sun.getDate()+6);
          var h=el('h4'); h.textContent='Week '+(i+1)+' ';
          h.appendChild(el('span')).textContent=fmtDay(k,{month:'short',day:'numeric'})+' to '+sun.toLocaleDateString('en-US',{month:'short',day:'numeric'})+' · '+weeks[k].length+(weeks[k].length>1?' posts':' post');
          sec.appendChild(h);
          var g=el('div','og-grid'); weeks[k].forEach(function(s){ g.appendChild(card(s)); });
          sec.appendChild(g); pn.appendChild(sec);
        });
      } else {
        var g=el('div','og-grid'); list.forEach(function(s){ g.appendChild(card(s)); }); pn.appendChild(g);
      }
      wrap.appendChild(pnOuter);
    });
  }
  /* Download links go through the viewer's save prompt (downloads capability); hidden where it isn't served */
  var dlLinks=[].slice.call(document.querySelectorAll('.og-post a[download]'));
  dlLinks.forEach(function(a){ a.hidden=true; });
  if(window.claude&&claude.use){ claude.use('downloads').then(function(dl){
    if(!dl) return;
    dlLinks.forEach(function(a){
      a.hidden=false;
      a.addEventListener('click',function(ev){
        ev.preventDefault();
        var name=a.getAttribute('href').split('/').pop(), lab=a.textContent;
        a.textContent='Preparing…';
        fetch(a.href).then(function(r){ if(!r.ok) throw new Error(r.status); return r.blob(); })
          .then(function(b){ return dl.save({filename:'Katapult-'+name, data:b}); })
          .then(function(){ a.textContent=lab; },function(){ a.textContent=lab; });
      });
    });
  }); }
  function setMonth(mid){
    var ps=[].slice.call(document.querySelectorAll('.og-mpanel'));
    if(!ps.length) return;
    if(!ps.some(function(p){return p.dataset.month===mid;})) mid=ps.map(function(p){return p.dataset.month;}).sort().reverse()[0];
    ps.forEach(function(p){ p.hidden=p.dataset.month!==mid; });
    [].forEach.call(document.querySelectorAll('.og-mbtn'),function(b){ b.setAttribute('aria-pressed',String(b.dataset.month===mid)); });
  }
"""

def patch_nav(src, months):
    a, b = src.index("<script>\n(function(){\n  /* Card index + hash router."), None
    nav = src[a:]
    end = nav.index("</script>")
    js = nav[:end]
    js2 = js.replace(
        "var POSTS=[].slice.call(document.querySelectorAll('main[data-ch=\"organic\"] > section.q4-post'));",
        "var POSTS=[].slice.call(document.querySelectorAll('main[data-ch=\"organic\"] > section.og-post'));\n"
        "  var ORG_MONTHS=" + json.dumps({m["id"]: {"label": m["label"], "status": m["status"], "lede": m["lede"], "note": m["status_note"]} for m in months}) + ";")
    if js2 == js:
        assert "section.og-post" in js, "POSTS line not found"
        # already patched: refresh the month data
        meta = json.dumps({m["id"]: {"label": m["label"], "status": m["status"], "lede": m["lede"], "note": m["status_note"]} for m in months})
        js2, n = re.subn(r"  var ORG_MONTHS=.*?;\n", lambda _: "  var ORG_MONTHS=" + meta + ";\n", js, count=1)
        assert n == 1, "ORG_MONTHS not found"
    js = js2
    s, e = js.index("  function buildPosts(){"), js.index("  /* ---- breadcrumbs")
    js = js[:s] + ORG_JS.lstrip("\n") + "\n" + js[e:]
    js = re.sub(r"  var postIds=POSTS\.map\(.*?\n  POSTS\.forEach\(function\(s,i\)\{ s\.insertBefore\(crumb\(.*?\n",
                "  var byMonth={}; POSTS.forEach(function(s){ (byMonth[s.dataset.month]=byMonth[s.dataset.month]||[]).push(s.id); });\n"
                "  POSTS.forEach(function(s){ var l=byMonth[s.dataset.month], mm=ORG_MONTHS[s.dataset.month]||{label:''}; "
                "s.insertBefore(crumb('#organic-'+s.dataset.month,'All '+mm.label+' posts',l,l.indexOf(s.id)),s.firstChild); });\n",
                js, flags=re.S)
    assert "byMonth" in js, "crumb line not found"
    js = js.replace("var ps=target.closest('section.q4-post');", "var ps=target.closest('section.og-post');")
    js = js.replace("if(hash==='organic'){ ch='organic'; view=ixO; }",
                    "if(hash==='organic'||hash.indexOf('organic-')===0){ ch='organic'; view=ixO; setMonth(hash.slice(8)); }")
    assert "setMonth(hash.slice(8))" in js and "section.og-post');" in js
    return src[:a] + js + nav[end:]

def build(hub_in, hub_out, files_dir):
    src = pathlib.Path(hub_in).read_text(encoding="utf-8")
    months = [month_nov(), month_oct()]
    stage(months, files_dir)
    posts_html = "\n".join(post_section(m, p) for m in months for p in m["posts"])
    org = f'<!-- ORG:START -->\n<main data-ch="organic">\n{posts_html}\n</main>\n<!-- ORG:END -->'
    if "<!-- ORG:START -->" in src:
        src = re.sub(r"<!-- ORG:START -->.*?<!-- ORG:END -->", lambda _: org, src, flags=re.S)
    else:
        s = src.index('<div class="q4-bar" data-ch="organic">'); e = src.index("</main>", s) + len("</main>")
        src = src[:s] + org + src[e:]
    if "/* ORG:START */" in src:
        src = re.sub(r"\n/\* ORG:START \*/.*?/\* ORG:END \*/\n", lambda _: ORG_CSS, src, flags=re.S)
    else:
        src = src.replace("/* NAV:END */", "/* NAV:END */" + ORG_CSS, 1)
    src = patch_nav(src, months)
    # masthead + tabs + organic index copy
    total = sum(len(m["posts"]) for m in months)
    frames = sum(len(p["slides"]) for m in months for p in m["posts"])
    reels = sum(1 for m in months for p in m["posts"] if p["reel"])
    src = re.sub(r'(<button type="button" class="hub-tab" role="tab" data-go="organic"[^>]*>Organic Social <small>)[^<]*(</small>)',
                 rf"\g<1>Oct + Nov · {total} posts\g<2>", src)
    mast = re.search(r'<div class="hub-ch-organic" data-ch="organic">.*?\n    </div>\n    </div>', src, re.S)
    if mast:
        notes = ('<div class="hub-ch-organic" data-ch="organic">\n    <div class="q4-notes">\n'
                 '      <p><b>October</b> is the final set, exported from Figma (page “kat os”) at native size. Click a slide for the full-size file.</p>\n'
                 '      <p><b>November</b> is built in the look of the October finals: sentence-case headlines with one Pink word, logo on covers and closers only, White cards on a solid plate, the Pink line and dots as the one accent. No CTA pill or disclaimer on the art.</p>\n'
                 '      <p><b>Copy</b> is the client\'s, placed as written. Every change (typos, dashes, emoji) is listed in the post\'s Notes.</p>\n'
                 '    </div>\n    </div>')
        src = src[:mast.start()] + notes + src[mast.end():]
    src = src.replace("<p>Fifteen posts for Q4 2026, grouped by theme. Open one to see all its slides at full size.</p>",
                      "<p>One tab per month. Open a post to see every slide at full size, the caption and the notes.</p>")
    src = src.replace('<div class="hub-ix-counts" id="hub-organic-counts"></div>\n', "")
    src = src.replace("<summary>How to read the slides</summary>", "<summary>About these posts</summary>")
    # the old October slide viewer looks for the removed safe-zone toggle: let it stand down
    src = src.replace("  function applySafe(){ app.classList.toggle('show-safe', t.checked); }\n  t.addEventListener('change', applySafe); applySafe();",
                      "  if(!t) return;\n  function applySafe(){ app.classList.toggle('show-safe', t.checked); }\n  t.addEventListener('change', applySafe); applySafe();")
    src = src.replace("<p>Pick a month. November opens on its posting calendar; October is the final set from Figma. Open a post for every slide at full size, the caption and the notes.</p>",
                      "<p>One tab per month. Open a post to see every slide at full size, the caption and the notes.</p>")
    src = src.replace("var c=mastOrg.querySelector('.q4-counts'); if(c) document.getElementById('hub-organic-counts').appendChild(c);", "")
    pathlib.Path(hub_out).write_text(src, encoding="utf-8")
    print(f"posts {total}, frames {frames}, reels {reels}")

if __name__ == "__main__":
    build(*sys.argv[1:4])
